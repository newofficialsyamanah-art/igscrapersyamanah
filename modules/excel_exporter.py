import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.cell_style import CellStyle
from openpyxl.utils import get_column_letter

# Fix openpyxl CellStyle bug: ensure applyFont, applyFill, applyBorder are written to cellXfs
# so Microsoft Excel, LibreOffice, and WPS Office properly render fills and font colors!
if "applyFill" not in CellStyle.__attrs__:
    CellStyle.__attrs__ = (
        'numFmtId', 'fontId', 'fillId', 'borderId',
        'applyFont', 'applyFill', 'applyBorder',
        'applyAlignment', 'applyProtection', 'pivotButton', 'quotePrefix', 'xfId'
    )
    _orig_from_array = CellStyle.from_array
    def _patched_from_array(cls, style):
        cs = _orig_from_array(style)
        if style.fontId > 0:
            cs.applyFont = True
        if style.fillId > 0:
            cs.applyFill = True
        if style.borderId > 0:
            cs.applyBorder = True
        return cs
    CellStyle.from_array = classmethod(_patched_from_array)

def generate_excel_data(users, target_username="data", scrape_type="followers"):
    wb = Workbook()
    ws = wb.active
    type_name = "Followers" if scrape_type == "followers" else "Following"
    ws.title = f"{type_name} @{target_username}"[:31]  # Excel sheet title max 31 chars

    # Column Headers on ROW 1
    headers = [
        "No",
        "Username",
        "Nama Lengkap",
        "Kategori / Status",
        "Followers",
        "Following",
        "Email Publik",
        "Status Akun",
        "Centang Biru",
        "Target Sumber",
        "Link Profil",
        "Waktu Scrape"
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    header_fill = PatternFill(start_color="FF1E293B", end_color="FF1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style='thin', color='FFCBD5E1'),
        right=Side(style='thin', color='FFCBD5E1'),
        top=Side(style='thin', color='FFCBD5E1'),
        bottom=Side(style='thin', color='FFCBD5E1')
    )

    # Style Row 1 Headers
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Pin Header Row on scroll (Freeze Panes)
    ws.freeze_panes = "A2"

    # Data Rows starting at Row 2
    row_idx = 2
    for idx, u in enumerate(users, 1):
        username = str(u.get("username", "")).lstrip("@")
        full_name = str(u.get("full_name") or "-")
        category = str(u.get("category_type") or "Follower Biasa")

        # Formatted follower & following counts
        f_val = u.get("follower_count_str")
        if not f_val or f_val in ["0", "-"]:
            raw_f = u.get("follower_count") or u.get("followers_count")
            f_val = f"{int(raw_f):,}" if raw_f else "-"
        
        g_val = u.get("following_count_str")
        if not g_val or g_val in ["0", "-"]:
            raw_g = u.get("following_count")
            g_val = f"{int(raw_g):,}" if raw_g else "-"

        email = str(u.get("email") or "-")
        status_akun = "Privat" if u.get("is_private") else "Publik"
        verified = "Ya" if u.get("is_verified") else "Tidak"
        sumber = str(u.get("target_source") or target_username).lstrip("@")
        link = f"https://www.instagram.com/{username}/"
        scraped_at = str(u.get("scraped_at") or "-")

        row = [
            idx,
            f"@{username}",
            full_name,
            category,
            f_val,
            g_val,
            email,
            status_akun,
            verified,
            f"@{sumber}",
            link,
            scraped_at
        ]
        ws.append(row)
        ws.row_dimensions[row_idx].height = 20

        # Style data cells
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border

            # Column specific alignments
            if col_idx in [1, 5, 6, 8, 9, 12]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Clickable profile link
            if col_idx == 11:
                cell.hyperlink = link
                cell.font = Font(name="Calibri", size=10, color="FF0284C7", underline="single")

            # Highlight email cell
            if col_idx == 7 and email != "-":
                cell.fill = PatternFill(start_color="FFDCFCE7", end_color="FFDCFCE7", fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color="FF166534")
            # Zebra striping
            elif row_idx % 2 == 1:
                cell.fill = PatternFill(start_color="FFF8FAFC", end_color="FFF8FAFC", fill_type="solid")

        row_idx += 1

    # Enable native Excel AutoFilter on all headers
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{max(row_idx - 1, 1)}"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
